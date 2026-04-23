import argparse
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_BASE_HEADERS = ["appId", "简体中文", "简体中文"]


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--lang-header", required=True)
    parser.add_argument("--entries-file", required=True)
    return parser.parse_args()


def load_entries(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise ValueError("entries-file must be a JSON array")

    entries = []
    for idx, item in enumerate(data, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"entry #{idx} must be an object")
        app_id = item.get("appId")
        zh = item.get("zh")
        target = item.get("target")
        if not all(isinstance(v, str) and v for v in (app_id, zh, target)):
            raise ValueError(f"entry #{idx} must include non-empty string appId, zh, target")
        entries.append((app_id, zh, target))
    return entries


def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.number_format = src.number_format
    if ws.row_dimensions[src_row].height is not None:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def validate_headers(ws, lang_header):
    headers = [ws.cell(1, col).value for col in range(1, 5)]
    expected = EXPECTED_BASE_HEADERS + [lang_header]
    if headers != expected:
        raise ValueError(
            f"unexpected headers in {ws.title}: {headers}, expected {expected}"
        )


def build_index(ws):
    index = {}
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if value:
            index[str(value)] = row
    return index


def upsert_workbook(source, output, sheet_name, lang_header, entries):
    wb = load_workbook(source)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"sheet not found: {sheet_name}")

    ws = wb[sheet_name]
    validate_headers(ws, lang_header)

    style_row = ws.max_row
    index = build_index(ws)
    added = 0
    updated = 0

    for app_id, zh, target in entries:
        if app_id in index:
            row = index[app_id]
            ws.cell(row, 2).value = zh
            ws.cell(row, 3).value = zh
            ws.cell(row, 4).value = target
            updated += 1
            continue

        row = ws.max_row + 1
        copy_row_style(ws, style_row, row, ws.max_column)
        ws.cell(row, 1).value = app_id
        ws.cell(row, 2).value = zh
        ws.cell(row, 3).value = zh
        ws.cell(row, 4).value = target
        index[app_id] = row
        added += 1

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {
        "output": str(output_path),
        "sheet": sheet_name,
        "langHeader": lang_header,
        "added": added,
        "updated": updated,
        "totalRows": ws.max_row,
    }


def main():
    args = parse_args()
    entries = load_entries(args.entries_file)
    result = upsert_workbook(
        source=args.source,
        output=args.output,
        sheet_name=args.sheet,
        lang_header=args.lang_header,
        entries=entries,
    )
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
